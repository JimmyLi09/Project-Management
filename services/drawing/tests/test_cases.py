"""历史案例：按公司统计表的列位读取两张 LED 工作表。夹具在测试里现场生成，
不提交真实统计表（其中有客户与联系人信息）。"""
import json

import openpyxl

from avdrawing.cases import main, pitch_of, read


def _row(width, **cells):
    r = [None] * width
    for col, v in cells.items():
        r[int(col[1:]) - 1] = v
    return r


def _book(tmp_path, sheets=("LED ongoing project", "LED completed project")):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if "LED ongoing project" in sheets:
        ws = wb.create_sheet("LED ongoing project")
        ws.append(["LED ONGOING"]); ws.append(["YEAR", "PROJECT NAME"]); ws.append([])
        ws.append(_row(21, c1=2026, c2="021-Nafa LED", c3="NAFA", c4="80 Bencoolen St", c5="+65 9000 0000",
                       c10="12,000", c11=3000, c12=36, c15=144, c16="P2.5 indoor", c17="18 KW",
                       c18="4 nos", c19="3+1 spare", c21="phase 1"))
        ws.append(_row(21, c1=2026))          # 没有项目名：跳过
    if "LED completed project" in sheets:
        ws = wb.create_sheet("LED completed project")
        ws.append(["LED COMPLETED"]); ws.append(["", "NO.", "NAME"]); ws.append(["FINISHED PROJECTS"])
        ws.append(_row(22, c2="128", c3="  C&K   T1 LED ", c4="C&K", c11=1350, c12=1350, c13=1.8225,
                       c16=12, c17="P 1.56 COB", c18=0.91, c20="1"))
        ws.append(_row(22, c17="P2.5", c11=2000, c12=1000, c13=2))          # 同项目的第二块屏
        ws.append(_row(22, c2="129", c3="Mock-up", c17="Samsung IER"))
        # 从进行中表整行搬来的行：A 列年份、B 列项目名、E 列联系人
        ws.append(_row(21, c1=2024, c2="112-Marina View", c3="IOI", c4="21 Park St", c5="Koo @ 8383 4325",
                       c10=12800, c11=4960, c12=63.488, c15=1240, c16="P2.5 Curve Wall", c17=29))
        ws.append(_row(21, c10=5120, c11=2160, c12=11.0592, c16="P3"))     # 其续行，没有年份
    wb.create_sheet("PROJECTOR").append(["不读取"])
    path = tmp_path / "links.xlsx"
    wb.save(path)
    return path


def test_两张工作表按列位读取(tmp_path):
    out = read(_book(tmp_path))
    assert out["sheets"] == {"LED ongoing project": 1, "LED completed project": 5}
    ongoing, done, done2, mock, moved, moved2 = out["cases"]
    assert ongoing == {
        "sourceSheet": "LED ongoing project", "status": "ongoing", "refNo": None, "year": 2026,
        "name": "021-Nafa LED", "client": "NAFA", "address": "80 Bencoolen St",
        "widthMm": 12000.0, "heightMm": 3000.0, "sqm": 36.0, "pitch": 2.5, "modules": 144, "kw": 18.0,
        "powerCable": "4 nos", "dataCable": "3+1 spare", "product": "P2.5 indoor", "remarks": "phase 1",
    }
    assert "+65" not in json.dumps(out), "MAIN CON（联系人电话）不读取"
    assert (done["status"], done["refNo"], done["name"], done["pitch"], done["sqm"]) == ("completed", "128", "C&K T1 LED", 1.56, 1.8225)
    assert (mock["pitch"], mock["sqm"], mock["modules"]) == (None, None, None)
    assert (done2["name"], done2["refNo"], done2["client"], done2["pitch"], done2["sqm"]) == ("C&K T1 LED", "128", "C&K", 2.5, 2.0)


def test_完成表里的进行中列位行与续行(tmp_path):
    *_, moved, moved2 = read(_book(tmp_path))["cases"]
    assert (moved["status"], moved["year"], moved["name"], moved["client"], moved["address"]) == ("completed", 2024, "112-Marina View", "IOI", "21 Park St")
    assert (moved["sqm"], moved["pitch"], moved["modules"], moved["kw"]) == (63.488, 2.5, 1240, 29.0)
    assert (moved2["name"], moved2["year"], moved2["sqm"], moved2["pitch"]) == ("112-Marina View", 2024, 11.0592, 3.0)


def test_点间距取自型号():
    assert [pitch_of(t) for t in ("P1.56 COB", "p 3.91 outdoor", "P10", "Samsung IER", None)] == [1.56, 3.91, 10.0, None, None]


def test_命令行_缺工作表与读不了的文件都给出错误(tmp_path, capsys):
    assert main([str(_book(tmp_path))]) == 0
    assert len(json.loads(capsys.readouterr().out)["cases"]) == 6
    only = _book(tmp_path, sheets=())
    assert main([str(only)]) == 2
    assert "没有找到 LED 工作表" in json.loads(capsys.readouterr().out)["error"]
    bad = tmp_path / "x.xlsx"
    bad.write_bytes(b"not a workbook")
    assert main([str(bad)]) == 2
    assert "无法读取统计表" in json.loads(capsys.readouterr().out)["error"]
