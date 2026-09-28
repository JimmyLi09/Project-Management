"""技术方案书：数值来自 payload，叙述随校验结果变化，表格状态按库内 / 库外。"""
from docx import Document

from avdrawing.proposal import narrative, render

P144 = {
    "title": "144 Chuan Grove", "client": "海晟置业", "packVersion": "led@1.0", "profile": "室内固装",
    "L": 4480, "H": 2560, "pitch": 2, "sqm": 11.4688, "pxW": 2240, "pxH": 1280, "mods": 224, "cabinets": 35,
    "kw": 5.7344, "nCircuit": 3, "nPowerCable": 4, "nDataRun": 6, "nDataCable": 7,
    "bom": [{"w": 640, "h": 480, "count": 28, "inLib": True}, {"w": 640, "h": 640, "count": 7, "inLib": True}],
    "custom": False, "circuitKw": 2.5, "circuitsW": [2458, 1638, 1638],
    "control": {"name": "诺瓦", "dataPx": 560000}, "viewMin": 3, "pwrDist": 18, "findings": [],
}


def _text(path):
    d = Document(path)
    paras = [p.text for p in d.paragraphs]
    cells = [[c.text for c in r.cells] for t in d.tables for r in t.rows]
    return paras, cells


def test_数值与段落(tmp_path):
    out = render(P144, tmp_path / "p.docx")
    paras, cells = _text(out)
    body = "\n".join(paras)
    assert "LED 显示屏系统技术方案" in paras[0]
    assert "144 Chuan Grove　|　海晟置业" in paras[1]
    assert "4480 × 2560 mm" in body and "P2 室内固装" in body and "2240 × 1280" in body
    assert "35 只箱体" in body and "224 块模组" in body and "无需定制" in body
    assert "3 个供电回路，另预留 1 路备用，合计电源线 4 根" in body
    assert "560,000 像素" in body and "需数据线 6 条，另预留 1 条备用，合计 7 根" in body
    assert "与 P2 的点间距配置相匹配" in body
    assert "强电井" not in body
    assert "规则包 led@1.0" in body
    assert ["640 × 480", "28 只", "库内标准"] in cells
    assert ["回路分配", "2458W / 1638W / 1638W", "按列均衡"] in cells
    assert "六、待确认事项" not in body


def test_没有客户名时不留分隔符(tmp_path):
    paras, _ = _text(render({**P144, "client": ""}, tmp_path / "p.docx"))
    assert paras[1] == "144 Chuan Grove"


def test_叙述跟随校验结果():
    p = {**P144, "viewMin": 1.5, "pwrDist": 35, "findings": [
        {"code": "LED-VD-01", "severity": "warn", "message": "近距离会看到颗粒"},
        {"code": "LED-PWR-07", "severity": "info", "message": "强电井距离 35 m"}]}
    text = "\n".join(narrative(p))
    assert "小于 P2 的推荐最小观看距离" in text
    assert "强电井距屏体 35 m" in text


def test_库外箱体与待确认事项(tmp_path):
    p = {**P144, "custom": True, "bom": [{"w": 640, "h": 480, "count": 28, "inLib": True},
                                          {"w": 640, "h": 400, "count": 7, "inLib": False}],
         "findings": [{"code": "LED-CAB-01", "severity": "warn", "message": "排布含库外规格"}]}
    paras, cells = _text(render(p, tmp_path / "p.docx"))
    body = "\n".join(paras)
    assert "需向厂家定制" in body
    assert ["640 × 400", "7 只", "库外，需定制"] in cells
    assert "六、待确认事项" in body and "[LED-CAB-01] 排布含库外规格" in body
