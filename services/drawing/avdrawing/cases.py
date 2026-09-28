"""历史案例：读取公司《All the Project Links 统计表》的两张 LED 工作表。

移植自 avcost-phase1 的 scripts/import_cases.py（2026-09-28）。只读、只解析，结果以
JSON 输出，由平台写入数据库。每一行是一块屏。相对原版：
  · 按行识别列位：已完成表里有一批行是从进行中表整行搬过来的（A 列是年份、B 列是
    项目名），原版按完成表列位读，项目名、客户、地址会错一列（地址栏读进联系人电话）；
  · 同一项目的第二、三块屏写在下一行、不再填项目名，原版跳过，这里并入上一个项目；
  · 进行中表也读取年份、客户与地址；完成表的表内编号一并保留；
  · 不读取 MAIN CON 列（其中是联系人电话）。

用法：python -m avdrawing.cases 统计表.xlsx
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl

# 列号从 1 开始。两种列位只差一列：完成表多一列 NO.，进行中表 A 列是年份。
LAYOUTS = {
    "ongoing": {"year": 1, "name": 2, "client": 3, "address": 4, "L": 10, "H": 11, "sqm": 12,
                "modules": 15, "type": 16, "kw": 17, "power_cable": 18, "data_cable": 19, "remarks": 21},
    "completed": {"ref_no": 2, "name": 3, "client": 4, "address": 5, "L": 11, "H": 12, "sqm": 13,
                  "modules": 16, "type": 17, "kw": 18, "power_cable": 19, "data_cable": 20, "remarks": 22},
}
SHEETS = {"LED ongoing project": "ongoing", "LED completed project": "completed"}
FIRST_ROW = 4
PROJECT = ("refNo", "year", "name", "client", "address")   # 续行（同项目的下一块屏）沿用上一行
SCREEN = ("L", "H", "sqm", "type")                         # 有其一才算一块屏


def num(v) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        m = re.search(r"-?\d+(?:\.\d+)?", v.replace(",", ""))
        if m:
            return float(m.group())
    return None


def text(v) -> str | None:
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def pitch_of(type_text) -> float | None:
    m = re.search(r"P\s*([0-9]+(?:\.[0-9]+)?)", str(type_text or ""), re.I)
    return float(m.group(1)) if m else None


def _year(v) -> int | None:
    y = num(v) if not hasattr(v, "year") else None
    return int(y) if y and 1990 < y < 2100 else None


def _cell(row, layout: str, key: str):
    i = LAYOUTS[layout].get(key)
    return row[i - 1] if i and i - 1 < len(row) else None


def _fits(row, layout: str) -> bool:
    """长 × 高 与面积对得上（±5%），说明这一行是这种列位。"""
    L, H, S = (num(_cell(row, layout, k)) for k in ("L", "H", "sqm"))
    return bool(L and H and S) and abs(L * H / 1e6 - S) <= 0.05 * S + 0.01


def _layout(row, status: str, prev: str) -> str:
    """A 列是年份 → 进行中列位；否则看尺寸对得上哪种；再看本表列位下有没有项目名；
    都判断不了（多为同项目的续行）就沿用上一行。"""
    if _year(_cell(row, "ongoing", "year")):
        return "ongoing"
    for layout in (status, *(k for k in LAYOUTS if k != status)):
        if _fits(row, layout):
            return layout
    return status if text(_cell(row, status, "name")) else prev


def read(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    cases, sheets = [], {}
    for sheet, status in SHEETS.items():
        if sheet not in wb.sheetnames:
            sheets[sheet] = None
            continue
        n, head, layout = 0, None, status
        for row in wb[sheet].iter_rows(min_row=FIRST_ROW, values_only=True):
            layout = _layout(row, status, layout)
            get = lambda k: _cell(row, layout, k)  # noqa: E731
            if not text(get("name")) and not (head and any(get(k) is not None for k in SCREEN)):
                continue
            modules = num(get("modules"))
            case = {
                "sourceSheet": sheet,
                "status": status,
                "refNo": text(get("ref_no")),
                "year": _year(get("year")),
                "name": text(get("name")),
                "client": text(get("client")),
                "address": text(get("address")),
                "widthMm": num(get("L")),
                "heightMm": num(get("H")),
                "sqm": num(get("sqm")),
                "pitch": pitch_of(get("type")),
                "modules": int(modules) if modules else None,
                "kw": num(get("kw")),
                "powerCable": text(get("power_cable")),
                "dataCable": text(get("data_cable")),
                "product": text(get("type")),
                "remarks": text(get("remarks")),
            }
            if case["name"]:
                head = case
            else:
                for k in PROJECT:
                    case[k] = case[k] or head[k]
            cases.append(case)
            n += 1
        sheets[sheet] = n
    return {"cases": cases, "sheets": sheets}


def main(argv: list[str]) -> int:
    if len(argv) != 1:
        print(__doc__, file=sys.stderr)
        return 2
    try:
        out = read(Path(argv[0]))
    except Exception as e:  # noqa: BLE001 — any unreadable workbook is reported, not crashed on
        print(json.dumps({"error": f"无法读取统计表：{e}"}, ensure_ascii=False))
        return 2
    if not out["cases"]:
        print(json.dumps({"error": "没有找到 LED 工作表或其中没有项目行，请检查文件与工作表名"}, ensure_ascii=False))
        return 2
    print(json.dumps(out, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
