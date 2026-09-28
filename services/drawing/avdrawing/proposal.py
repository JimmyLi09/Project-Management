"""LED 技术方案书（Word）排版。

内容（章节、段落、表格、待确认事项、版权页文字，中文或英文）全部由 TypeScript 内核
src/av/core/proposal.ts 的 ProposalDoc 给出，这里只负责排版，不产生任何数字或叙述：
  · 封面：AUDAX 标志、标题、项目 / 客户 / 日期；
  · 版权页；
  · 正文各章；
  · 除封面外每页页眉为标志 + AUDAX，页脚为文档名与页码。
模板结构移植自 avcost-phase1 的 proposal.py（2026-09-28）。

用法：python -m avdrawing.proposal doc.json out.docx
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ASSETS = Path(__file__).resolve().parent / "assets"
LOGO = ASSETS / "audax-logo.png"   # 标志 + 字标，字标已改为深色，适合白纸
MARK = ASSETS / "audax-mark.png"   # 仅标志

NAVY = "16293B"
HEAD = "1F3864"
GREY = "5B6168"
SEVERITY = {"block": "7A2E2E", "warn": "8A4419", "info": "3D5A73"}
#: (西文字体, 东亚字体)；英文版正文用 Calibri，个别中文（如客户名）仍落在雅黑
FONTS = {"zh": ("Microsoft YaHei", "Microsoft YaHei"), "en": ("Calibri", "Microsoft YaHei")}


class _Style:
    def __init__(self, lang: str):
        self.latin, self.east = FONTS.get(lang, FONTS["zh"])

    def run(self, run, size, bold=False, color=None):
        run.font.name = self.latin
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), self.east)
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        return run

    def para(self, container, text="", size=10.5, bold=False, color=None, after=6, align=None, before=0):
        p = container.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        if align is not None:
            p.alignment = align
        if text:
            self.run(p.add_run(text), size, bold, color)
        return p


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _bottom_border(paragraph, color=NAVY):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", color)):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    ppr.append(bdr)


def _page_field(paragraph, st: "_Style", size, color):
    """A PAGE field whose shown number carries the given size and colour."""
    def fld(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    for part in (fld("begin"), instr, fld("separate"), "1", fld("end")):
        run = st.run(paragraph.add_run("" if not isinstance(part, str) else part), size, color=color)
        if not isinstance(part, str):
            run._r.append(part)


def _shade(cell, fill="E8EEF3"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _letterhead(doc, section, st: _Style, doc_title: str, footer: str):
    section.different_first_page_header_footer = True   # the cover carries the full logo instead
    head = section.header.paragraphs[0]
    head.style = doc.styles["Normal"]   # the built-in Header style brings its own centre tab stop
    head.paragraph_format.tab_stops.add_tab_stop(section.page_width - section.left_margin - section.right_margin,
                                                 WD_TAB_ALIGNMENT.RIGHT)
    head.add_run().add_picture(str(MARK), height=Cm(0.75))
    st.run(head.add_run("  AUDAX"), 13, bold=True, color=NAVY)
    st.run(head.add_run("\t" + doc_title), 8.5, color=GREY)
    _bottom_border(head)

    foot = section.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.run(foot.add_run(footer + "  ·  "), 8, color=GREY)
    _page_field(foot, st, 8, GREY)


def _cover(doc, st: _Style, d: dict):
    st.para(doc, after=110)
    logo = st.para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=36)
    logo.add_run().add_picture(str(LOGO), width=Cm(7))
    st.para(doc, d["title"], size=24, bold=True, color=HEAD, align=WD_ALIGN_PARAGRAPH.CENTER, after=40)
    for row in d["cover"]:
        p = st.para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
        st.run(p.add_run(f"{row['label']}  "), 11, color=GREY)
        st.run(p.add_run(row["value"]), 12, bold=True, color=NAVY)
    _page_break(doc)


def _copyright(doc, st: _Style, d: dict):
    st.para(doc, after=300)
    st.para(doc, d["copyright"]["heading"], size=13, bold=True, color=HEAD, after=10)
    for t in d["copyright"]["paragraphs"]:
        st.para(doc, t, size=9.5, color=GREY, after=8)
    _page_break(doc)


def _table(doc, st: _Style, table: dict):
    t = doc.add_table(rows=1, cols=len(table["header"]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(table["header"]):
        cell = t.rows[0].cells[i]
        _shade(cell)
        st.run(cell.paragraphs[0].add_run(h), 9.5, bold=True, color=NAVY)
    for row in table["rows"]:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            st.run(cells[i].paragraphs[0].add_run(str(v)), 9.5)
    for row in t.rows:
        for cell in row.cells:
            fmt = cell.paragraphs[0].paragraph_format
            fmt.space_before, fmt.space_after = Pt(3), Pt(3)
    st.para(doc, after=4)


def build(d: dict) -> Document:
    st = _Style(d.get("lang", "zh"))
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = st.latin
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), st.east)
    normal.font.size = Pt(10.5)

    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)       # A4
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin, section.bottom_margin = Cm(2.4), Cm(2)
    _letterhead(doc, section, st, d["title"], d["footer"])

    _cover(doc, st, d)
    _copyright(doc, st, d)
    for s in d["sections"]:
        st.para(doc, s["heading"], size=15, bold=True, color=HEAD, after=8, before=6).paragraph_format.keep_with_next = True
        for text in s.get("paragraphs", []):
            st.para(doc, text, after=8)
        if s.get("table"):
            _table(doc, st, s["table"])
        for text in s.get("notes", []):
            st.para(doc, "· " + text, size=9.5, color=GREY, after=4)
        for it in s.get("items", []):
            st.para(doc, it["text"], size=9.5, color=SEVERITY.get(it["severity"], GREY), after=4)
    return doc


def render(doc_json: dict, out: Path) -> Path:
    build(doc_json).save(out)
    return out


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print(__doc__, file=sys.stderr)
        return 2
    src, dst = Path(argv[0]), Path(argv[1])
    render(json.loads(src.read_text(encoding="utf-8")), dst)
    print(dst)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
